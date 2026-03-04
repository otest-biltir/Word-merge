from __future__ import annotations

import shutil
from pathlib import Path
from typing import Iterable

from docx import Document
from docxcompose.composer import Composer

CATEGORY_KEYS = {
    "Pre photos": "pre",
    "Post photos": "post",
    "Teardown photos": "teardown",
    "Handle side cover": "handle_side_cover",
}

PLACEHOLDER_ALIASES = {
    1: ["{{IMG1}}", "[IMG1]", "<IMG1>", "{{PHOTO1}}"],
    2: ["{{IMG2}}", "[IMG2]", "<IMG2>", "{{PHOTO2}}"],
    3: ["{{IMG3}}", "[IMG3]", "<IMG3>", "{{PHOTO3}}"],
    4: ["{{IMG4}}", "[IMG4]", "<IMG4>", "{{PHOTO4}}"],
    5: ["{{IMG5}}", "[IMG5]", "<IMG5>", "{{PHOTO5}}"],
    6: ["{{IMG6}}", "[IMG6]", "<IMG6>", "{{PHOTO6}}"],
}


def chunks(values: list[Path], chunk_size: int) -> Iterable[list[Path]]:
    for i in range(0, len(values), chunk_size):
        yield values[i : i + chunk_size]


def _replace_placeholder_with_image(doc: Document, placeholder: str, image_path: Path) -> bool:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, "")
                run.add_picture(str(image_path))
                return True

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if placeholder in run.text:
                            run.text = run.text.replace(placeholder, "")
                            run.add_picture(str(image_path))
                            return True
    return False


def render_template_page(template_path: Path, image_chunk: list[Path], out_path: Path) -> Path:
    if not template_path.exists():
        raise FileNotFoundError(f"Template bulunamadı: {template_path}")

    shutil.copy(template_path, out_path)
    document = Document(str(out_path))

    for idx, image_path in enumerate(image_chunk, start=1):
        aliases = PLACEHOLDER_ALIASES[idx]
        inserted = any(_replace_placeholder_with_image(document, placeholder, image_path) for placeholder in aliases)
        if not inserted:
            raise ValueError(
                f"Template içinde {idx}. görsel için placeholder bulunamadı. "
                f"Beklenen örnek: {aliases[0]}"
            )

    document.save(str(out_path))
    return out_path


def merge_docx_files(docx_paths: list[Path], output_path: Path) -> Path:
    if not docx_paths:
        raise ValueError("Birleştirilecek dosya yok")

    master = Document(str(docx_paths[0]))
    composer = Composer(master)

    for path in docx_paths[1:]:
        composer.append(Document(str(path)))

    composer.save(str(output_path))
    return output_path


def default_output_filename(selected_categories: list[str]) -> str:
    names = [CATEGORY_KEYS[cat] for cat in selected_categories if cat in CATEGORY_KEYS]
    if not names:
        return "merged.docx"
    return f"{'_'.join(names)}.docx"
